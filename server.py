"""Local and cloud-ready server for Denta Atelier with SQLite database and export features."""
import io
import json
import os
import urllib.parse
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

import database

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False


class Handler(SimpleHTTPRequestHandler):
    def end_headers(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(200)
        self.end_headers()

    def send_json(self, data, status=200):
        body = json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path

        if path == "/api/data":
            try:
                data = database.get_all_data()
                self.send_json(data)
            except Exception as e:
                self.send_json({"error": str(e)}, status=500)
            return

        if path == "/api/status":
            self.send_json({
                "status": "online",
                "database_type": "sqlite",
                "database_path": database.get_db_path()
            })
            return

        if path == "/api/backup":
            try:
                data = database.get_all_data()
                content = json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")
                self.send_file(content, "application/json", "denta-atelier-backup.json")
            except Exception as e:
                self.send_json({"error": str(e)}, status=500)
            return

        # Default static file handling (index.html, etc.)
        super().do_GET()

    def do_POST(self):
        length = int(self.headers.get("Content-Length", 0))
        raw_body = self.rfile.read(length) if length > 0 else b"{}"
        try:
            payload = json.loads(raw_body.decode("utf-8") or "{}")
        except Exception:
            payload = {}

        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path

        if path in ("/api/data", "/api/save"):
            try:
                database.save_all_data(payload)
                self.send_json({"success": True, "message": "Datele au fost salvate în baza de date locală."})
            except Exception as e:
                self.send_json({"success": False, "error": str(e)}, status=500)
            return

        if path == "/api/reset":
            try:
                database.reset_to_seed()
                self.send_json({"success": True, "message": "Baza de date a fost resetată la valorile demo."})
            except Exception as e:
                self.send_json({"success": False, "error": str(e)}, status=500)
            return

        if path == "/export-visits":
            if not HAS_OPENPYXL:
                self.send_error(501, "openpyxl not installed")
                return
            stream = io.BytesIO()
            book = Workbook()
            sheet = book.active
            sheet.title = "Jurnal vizite"
            sheet.append(["Pacient", "Telefon", "Data", "Servicii"])
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1B1C20")
            for row in payload.get("rows", []):
                sheet.append(row)
            for column, width in {"A": 28, "B": 18, "C": 16, "D": 25}.items():
                sheet.column_dimensions[column].width = width
            book.save(stream)
            self.send_file(stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "jurnal-vizite.xlsx")
            return

        if path == "/export-pdf":
            if not HAS_REPORTLAB:
                self.send_error(501, "reportlab not installed")
                return
            stream = io.BytesIO()
            pdf = canvas.Canvas(stream, pagesize=A4)
            pdf.setTitle("Denta Atelier")
            y = 800
            pdf.setFont("Helvetica-Bold", 17)
            pdf.drawString(48, y, "Denta Atelier")
            y -= 34
            pdf.setFont("Helvetica", 10)
            for raw in str(payload.get("text", "Document Denta Atelier")).splitlines() or [""]:
                line = raw.encode("ascii", "replace").decode("ascii")
                while stringWidth(line, "Helvetica", 10) > 510:
                    cut = max(line.rfind(" ", 0, 90), 1)
                    pdf.drawString(48, y, line[:cut])
                    y -= 15
                    line = line[cut:].lstrip()
                pdf.drawString(48, y, line)
                y -= 15
                if y < 55:
                    pdf.showPage()
                    y = 800
                    pdf.setFont("Helvetica", 10)
            pdf.save()
            self.send_file(stream.getvalue(), "application/pdf", "denta-atelier-document.pdf")
            return

        self.send_error(404)

    def send_file(self, content, mime, name):
        self.send_response(200)
        self.send_header("Content-Type", mime)
        self.send_header("Content-Disposition", f'attachment; filename="{name}"')
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)


def run():
    port = int(os.environ.get("PORT", 8080))
    # When deployed in cloud containers (e.g. Render, Railway), bind to 0.0.0.0
    # When running locally without explicit PORT, 127.0.0.1 or 0.0.0.0 can be used
    host = os.environ.get("HOST", "0.0.0.0" if "PORT" in os.environ else "127.0.0.1")
    print(f"=====================================================")
    print(f"🦷 Denta Atelier rulează pe http://{host}:{port}")
    print(f"📦 Baza de date SQLite: {database.get_db_path()}")
    print(f"=====================================================")
    database.init_db()
    server = ThreadingHTTPServer((host, port), Handler)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nOprire server.")
        server.server_close()


if __name__ == "__main__":
    run()
