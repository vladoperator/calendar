import io
import json
from http.server import BaseHTTPRequestHandler
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


class handler(BaseHTTPRequestHandler):
    def do_POST(self):
        size = int(self.headers.get("Content-Length", 0))
        payload = json.loads(self.rfile.read(size) or b"{}")
        stream = io.BytesIO()
        book = Workbook()
        sheet = book.active
        sheet.title = "Jurnal vizite"
        sheet.append(["Pacient", "Telefon", "Data", "Servicii"])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1B1C20")
        for row in payload.get("rows", []):
            sheet.append(row)
        for column, width in {"A": 28, "B": 18, "C": 16, "D": 25}.items():
            sheet.column_dimensions[column].width = width
        book.save(stream)
        content = stream.getvalue()
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", 'attachment; filename="jurnal-vizite.xlsx"')
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)
